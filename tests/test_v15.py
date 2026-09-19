"""V1.5 的关键边界和报表对账测试。"""

import csv
import math
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
import config  # noqa: E402
import main  # noqa: E402


def row(**changes):
    value = {
        "达人名称": "测试达人", "粉丝数": "10000", "播放量": "1000", "点赞数": "50",
        "评论数": "20", "转发数": "10", "投放消耗": "100", "线索量": "10",
        "成交量": "2", "GMV": "200",
    }
    value.update(changes)
    return value


class ReportV15Tests(unittest.TestCase):
    def test_real_zero_denominators(self):
        data = main.analyze(pd.DataFrame([
            row(达人名称="播放为0", 播放量="0"),
            row(达人名称="线索为0", 线索量="0"),
            row(达人名称="消耗为0", 投放消耗="0"),
        ]))
        self.assertEqual(data.iloc[0]["播放量"], 0)
        self.assertTrue(math.isnan(data.iloc[0]["互动率"]))
        self.assertTrue(math.isnan(data.iloc[0]["CPM"]))
        self.assertTrue(math.isnan(data.iloc[1]["CPL"]))
        self.assertTrue(math.isnan(data.iloc[1]["线索转化率"]))
        self.assertTrue(math.isnan(data.iloc[2]["ROI"]))
        self.assertEqual(data.attrs["quality"]["zero_rows"], 3)

    def test_missing_format_and_negative_have_distinct_counts(self):
        data = main.analyze(pd.DataFrame([
            row(达人名称="缺失", GMV=""),
            row(达人名称="格式", GMV="待补录"),
            row(达人名称="负数", 投放消耗="-20"),
        ]))
        self.assertTrue(all(math.isnan(data.iloc[i][field]) for i, field in ((0, "GMV"), (1, "GMV"), (2, "投放消耗"))))
        quality = data.attrs["quality"]
        self.assertEqual(quality["fields"]["GMV"]["缺失值"], 1)
        self.assertEqual(quality["fields"]["GMV"]["格式异常"], 1)
        self.assertEqual(quality["fields"]["投放消耗"]["负数异常"], 1)
        self.assertEqual(quality["pending_rows"], 3)
        self.assertTrue(all(main.classify(item)[0] == "继续观察" for _, item in data.iterrows()))

    def test_currency_and_thousands_parse(self):
        data = main.analyze(pd.DataFrame([row(播放量="1,000", 投放消耗="¥1,234.50", GMV="￥2,000.00")]))
        self.assertEqual(data.iloc[0]["播放量"], 1000)
        self.assertEqual(data.iloc[0]["投放消耗"], 1234.5)
        self.assertEqual(data.iloc[0]["GMV"], 2000)
        self.assertAlmostEqual(data.iloc[0]["ROI"], 2000 / 1234.5)
        self.assertEqual(data.iloc[0]["数据异常说明"], "")

    def test_missing_name_is_pending_and_excluded_from_summary(self):
        data = main.analyze(pd.DataFrame([row(达人名称=""), row(达人名称="完整")]))
        labels = data.apply(main.classify, axis=1)
        values = main.calculate_overall(data, labels)
        self.assertEqual(data.attrs["quality"]["fields"]["达人名称"]["缺失值"], 1)
        self.assertEqual(labels.iloc[0][0], "继续观察")
        self.assertIn("待补数据", labels.iloc[0][1])
        self.assertEqual(values["有效汇总记录数"], 1)

    def test_csv_input(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.csv"
            with path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=main.FIELDS)
                writer.writeheader()
                writer.writerow(row())
            loaded = main.load_input(path)
            self.assertEqual(len(loaded), 1)
            self.assertEqual(main.analyze(loaded).iloc[0]["ROI"], 2)

    def test_xlsx_input(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.xlsx"
            workbook = Workbook()
            workbook.active.append(main.FIELDS)
            sample = row()
            workbook.active.append([sample[field] for field in main.FIELDS])
            workbook.save(path)
            loaded = main.load_input(path)
            self.assertEqual(len(loaded), 1)
            self.assertEqual(main.analyze(loaded).iloc[0]["ROI"], 2)

    def test_overall_reconciles_and_excludes_unknown(self):
        data = main.analyze(pd.DataFrame([
            row(达人名称="甲"),
            row(达人名称="乙", 播放量="2000", 点赞数="100", 评论数="40", 转发数="20", 投放消耗="200", 线索量="20", 成交量="4", GMV="400"),
            row(达人名称="待补", GMV=""),
        ]))
        values = main.calculate_overall(data, data.apply(main.classify, axis=1))
        self.assertEqual(values["有效汇总记录数"], 2)
        self.assertEqual(values["总投放消耗"], 300)
        self.assertEqual(values["总播放量"], 3000)
        self.assertEqual(values["总GMV"], 600)
        self.assertEqual(values["总互动数"], 240)
        self.assertEqual(values["整体ROI"], 2)
        self.assertEqual(values["整体CPM"], 100)
        self.assertEqual(values["整体CPL"], 10)
        self.assertAlmostEqual(values["整体互动率"], 0.08)
        self.assertAlmostEqual(values["整体线索转化率"], 0.2)

    def test_classification_explains_actual_trigger(self):
        data = main.analyze(pd.DataFrame([
            row(达人名称="复投"),
            row(达人名称="淘汰", 投放消耗="300", 线索量="1", GMV="30"),
            row(达人名称="观察", GMV="100"),
            row(达人名称="待补", GMV=""),
        ]))
        results = [main.classify(item) for _, item in data.iterrows()]
        self.assertEqual([item[0] for item in results], ["建议复投", "建议淘汰", "继续观察", "继续观察"])
        self.assertIn("ROI≥{:g}".format(config.REINVEST_MIN_ROI), results[0][2])
        self.assertIn("ROI<{:g}".format(config.ELIMINATE_BELOW_ROI), results[1][2])
        self.assertIn("CPL>{:g}".format(config.ELIMINATE_ABOVE_CPL), results[1][2])
        self.assertIn("待补数据", results[3][1])

    def test_workbook_charts_quality_config_and_png(self):
        data = main.analyze(pd.DataFrame([row(达人名称="复投"), row(达人名称="待补", GMV=""), row(达人名称="观察", GMV="100")]))
        with tempfile.TemporaryDirectory() as folder:
            report = Path(folder) / "report.xlsx"
            picture = Path(folder) / "dashboard.png"
            values, labels = main.create_report(data, report)
            main.save_dashboard_image(data, values, labels, picture)
            book = load_workbook(report)
            self.assertEqual(book.sheetnames, ["数据看板", "整体数据明细", "达人表现排名", "达人分层建议", "数据质量报告", "整体复盘", "参数配置说明"])
            charts = book["数据看板"]._charts
            self.assertEqual(len(charts), 3)
            self.assertEqual(charts[1].x_axis.scaling.orientation, "maxMin")
            self.assertEqual(charts[1].y_axis.scaling.orientation, "minMax")
            self.assertEqual(charts[1].y_axis.scaling.min, 0)
            self.assertEqual(charts[2].x_axis.numFmt.formatCode, "#,##0,")
            quality = book["数据质量报告"]
            self.assertEqual(quality["B2"].value, 3)
            self.assertEqual(quality["B5"].value, 1)
            self.assertIn("本页阈值仅用于个人项目功能演示", book["参数配置说明"]["A8"].value)
            self.assertEqual(book["达人分层建议"]["D1"].value, "触发规则")
            with Image.open(picture) as image:
                self.assertEqual(image.size, (2560, 1440))


if __name__ == "__main__":
    unittest.main()
