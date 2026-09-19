"""达人投放数据自动化复盘工具（V1.5）。

本项目所使用数据均为模拟测试数据，仅用于功能演示，不代表任何公司真实经营数据。
"""

import argparse
import csv
import math
import random
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from dashboard_image import save_dashboard_image


ROOT = Path(__file__).resolve().parents[1]
DISCLAIMER = "本项目所使用数据均为模拟测试数据，仅用于功能演示，不代表任何公司真实经营数据。"
FIELDS = ["达人名称", "粉丝数", "播放量", "点赞数", "评论数", "转发数", "投放消耗", "线索量", "成交量", "GMV"]
NUMBER_FIELDS = FIELDS[1:]
COUNT_FIELDS = set(FIELDS[1:6] + ["线索量", "成交量"])
MONEY_FORMAT = '"¥"#,##0.00;[Red]-"¥"#,##0.00'
PERCENT_FORMAT = "0.00%"
RATIO_FORMAT = "0.00"
COUNT_FORMAT = "#,##0"
QUALITY_KINDS = ("缺失值", "格式异常", "负数异常", "真实0")


def generate_sample(path):
    """用固定随机种子生成 100 条可复现的模拟测试数据。"""
    rng = random.Random(20260918)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = []
    for number in range(1, 101):
        tier = rng.choices(["small", "medium", "large"], weights=[60, 30, 10])[0]
        low, high = {"small": (5000, 100000), "medium": (100000, 600000), "large": (600000, 2000000)}[tier]
        fans = int(math.exp(rng.uniform(math.log(low), math.log(high))))
        views = int(fans * rng.uniform(0.25, 1.7))
        interactions = int(views * rng.uniform(0.018, 0.11))
        likes = int(interactions * rng.uniform(0.70, 0.84))
        comments = int(interactions * rng.uniform(0.07, 0.15))
        shares = interactions - likes - comments
        cost = round(views * rng.uniform(18, 95) / 1000, 2)
        leads = int(views * rng.uniform(0.0003, 0.004))
        sales = int(leads * rng.uniform(0.03, 0.20))
        gmv = round(sales * rng.uniform(100, 700), 2)
        rows.append({
            "达人名称": "模拟达人{:03d}".format(number),
            "粉丝数": fans, "播放量": views, "点赞数": likes, "评论数": comments,
            "转发数": shares, "投放消耗": cost, "线索量": leads,
            "成交量": sales, "GMV": gmv, "数据声明": DISCLAIMER,
        })

    # 边界样本用于演示分母为 0、缺失值与异常格式的处理。
    rows[5].update({"播放量": 0, "点赞数": 0, "评论数": 0, "转发数": 0, "线索量": 0, "成交量": 0, "GMV": 0})
    rows[24].update({"线索量": 0, "成交量": 0, "GMV": 0})
    rows[56].update({"投放消耗": 0})
    rows[44].update({"评论数": ""})
    rows[71].update({"投放消耗": -50})
    rows[88].update({"GMV": "待补录"})
    rows[9]["投放消耗"] = "¥{:,.2f}".format(rows[9]["投放消耗"])
    rows[10]["播放量"] = "{:,}".format(rows[10]["播放量"])

    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS + ["数据声明"])
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def load_input(path):
    if not path.is_file():
        raise ValueError("找不到输入文件：{}。可先运行 --generate-sample。".format(path))
    suffix = path.suffix.lower()
    if suffix == ".csv":
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                frame = pd.read_csv(path, dtype=str, keep_default_na=False, encoding=encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("CSV 编码无法识别，请另存为 UTF-8 或 GB18030。")
    elif suffix == ".xlsx":
        frame = pd.read_excel(path, dtype=str, keep_default_na=False, engine="openpyxl")
    else:
        raise ValueError("仅支持 .csv 和 .xlsx 输入文件。")
    frame.columns = [str(column).strip() for column in frame.columns]
    missing = [column for column in FIELDS if column not in frame.columns]
    if missing:
        raise ValueError("输入文件缺少字段：{}".format("、".join(missing)))
    if frame.empty:
        raise ValueError("输入文件没有数据行。")
    return frame


def clean_number(value, field):
    if pd.isna(value):
        return float("nan"), "{}缺失，待补数据".format(field), "缺失值"
    raw = str(value).strip()
    if not raw or raw.lower() in ("nan", "null", "none", "na", "n/a"):
        return float("nan"), "{}缺失，待补数据".format(field), "缺失值"
    normalized = re.sub(r"\s+", "", raw).replace(",", "").replace("，", "")
    normalized = normalized.replace("¥", "").replace("￥", "").replace("$", "")
    try:
        number = float(normalized)
    except ValueError:
        return float("nan"), "{}格式异常，待补数据".format(field), "格式异常"
    if not math.isfinite(number):
        return float("nan"), "{}格式异常，待补数据".format(field), "格式异常"
    if number < 0:
        return float("nan"), "{}非法负数，待补数据".format(field), "负数异常"
    if field in COUNT_FIELDS and not number.is_integer():
        return float("nan"), "{}非整数格式异常，待补数据".format(field), "格式异常"
    if number == 0:
        return int(number) if field in COUNT_FIELDS else number, None, "真实0"
    return int(number) if field in COUNT_FIELDS else number, None, None


def safe_divide(numerator, denominator):
    return numerator / denominator.replace(0, float("nan"))


def analyze(source):
    frame = source.copy()
    issues = [[] for _ in range(len(frame))]
    quality = {field: {kind: 0 for kind in QUALITY_KINDS} for field in FIELDS}
    zero_rows = set()
    pending_rows = set()
    for index, value in enumerate(frame["达人名称"]):
        if pd.isna(value) or not str(value).strip():
            frame.iat[index, frame.columns.get_loc("达人名称")] = "未命名达人_{}".format(index + 1)
            issues[index].append("达人名称缺失")
            quality["达人名称"]["缺失值"] += 1
            pending_rows.add(index)
        else:
            frame.iat[index, frame.columns.get_loc("达人名称")] = str(value).strip()
    for field in NUMBER_FIELDS:
        cleaned = []
        for index, value in enumerate(frame[field]):
            number, issue, kind = clean_number(value, field)
            cleaned.append(number)
            if kind:
                quality[field][kind] += 1
            if kind == "真实0":
                zero_rows.add(index)
            elif kind:
                pending_rows.add(index)
            if issue:
                issues[index].append(issue)
        frame[field] = cleaned

    frame["互动率"] = safe_divide(frame["点赞数"] + frame["评论数"] + frame["转发数"], frame["播放量"])
    frame["CPM"] = safe_divide(frame["投放消耗"], frame["播放量"]) * 1000
    frame["CPL"] = safe_divide(frame["投放消耗"], frame["线索量"])
    frame["线索转化率"] = safe_divide(frame["成交量"], frame["线索量"])
    frame["ROI"] = safe_divide(frame["GMV"], frame["投放消耗"])

    for index in range(len(frame)):
        row = frame.iloc[index]
        if row["播放量"] == 0:
            issues[index].append("播放量为0，互动率和CPM无法计算")
        if row["线索量"] == 0:
            issues[index].append("线索量为0，CPL和线索转化率无法计算")
        if row["投放消耗"] == 0:
            issues[index].append("投放消耗为0，ROI无法计算")
    frame["数据异常说明"] = ["；".join(item) for item in issues]
    frame.attrs["quality"] = {"fields": quality, "zero_rows": len(zero_rows), "pending_rows": len(pending_rows)}
    return frame


def classify(row):
    """返回分层结果、原因和实际触发的演示规则。"""
    issue = row["数据异常说明"]
    pending_fields = [field for field in NUMBER_FIELDS if pd.isna(row[field])]
    if "达人名称缺失" in issue:
        pending_fields.insert(0, "达人名称")
    if pending_fields:
        return "继续观察", "待补数据：{}需核对".format("、".join(pending_fields)), "关键字段缺失/异常 → 继续观察/待补数据"
    if row["播放量"] == 0 or row["投放消耗"] == 0:
        zero_fields = [field for field in ("播放量", "投放消耗") if row[field] == 0]
        return "继续观察", "{}为真实0，核心指标无法计算".format("、".join(zero_fields)), "播放量=0或投放消耗=0 → 继续观察"
    if row["ROI"] >= config.REINVEST_MIN_ROI and row["CPL"] <= config.REINVEST_MAX_CPL and row["线索量"] >= config.REINVEST_MIN_LEADS:
        reason = "ROI {:.2f}、CPL ¥{:.2f}、线索量 {} 均达标".format(row["ROI"], row["CPL"], int(row["线索量"]))
        rule = "ROI≥{:g} + CPL≤{:g} + 线索量≥{:g} → 建议复投".format(config.REINVEST_MIN_ROI, config.REINVEST_MAX_CPL, config.REINVEST_MIN_LEADS)
        return "建议复投", reason, rule
    causes = []
    triggers = []
    if row["ROI"] < config.ELIMINATE_BELOW_ROI:
        causes.append("ROI {:.2f}偏低".format(row["ROI"]))
        triggers.append("ROI<{:g}".format(config.ELIMINATE_BELOW_ROI))
    if pd.notna(row["CPL"]) and row["CPL"] > config.ELIMINATE_ABOVE_CPL:
        causes.append("CPL ¥{:.2f}偏高".format(row["CPL"]))
        triggers.append("CPL>{:g}".format(config.ELIMINATE_ABOVE_CPL))
    if row["线索量"] == 0:
        causes.append("线索量为0")
        triggers.append("线索量=0")
    if causes:
        return "建议淘汰", "；".join(causes), "、".join(triggers) + "（任一触发）→ 建议淘汰"
    return "继续观察", "未触发复投或淘汰阈值", "其余情况 → 继续观察"


def excel_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, (float, int)) and not math.isfinite(float(value)):
        return None
    return value


def append_frame(sheet, frame):
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        sheet.append([excel_value(value) for value in row])


def style_sheet(sheet, percent_columns=(), money_columns=(), count_columns=(), roi_column=None, issue_column=None):
    sheet.freeze_panes = "B2" if sheet.title != "整体复盘" else "A2"
    sheet.sheet_view.showGridLines = False
    if sheet.title != "整体复盘":
        sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="233B53")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="263238")
            cell.alignment = Alignment(vertical="center", horizontal="right" if isinstance(cell.value, (int, float)) else "left")
    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 101) + 1)]
        width = min(48, max(12, max((len(value) for value in values), default=0) * 1.3 + 2))
        sheet.column_dimensions[letter].width = width
    for column_name in percent_columns:
        column = find_column(sheet, column_name)
        for cells in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for cell in cells:
                cell.number_format = PERCENT_FORMAT
    for column_name in money_columns:
        column = find_column(sheet, column_name)
        for cells in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for cell in cells:
                cell.number_format = MONEY_FORMAT
    for column_name in count_columns:
        column = find_column(sheet, column_name)
        for cells in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for cell in cells:
                cell.number_format = COUNT_FORMAT
    if roi_column and sheet.max_row > 1:
        letter = get_column_letter(find_column(sheet, roi_column))
        green = PatternFill("solid", fgColor="DDF2E4")
        sheet.conditional_formatting.add("{}2:{}{}".format(letter, letter, sheet.max_row), CellIsRule(operator="greaterThanOrEqual", formula=[str(config.REINVEST_MIN_ROI)], fill=green))
        for cell in sheet[letter][1:]:
            cell.number_format = RATIO_FORMAT
    if issue_column and sheet.max_row > 1:
        letter = get_column_letter(find_column(sheet, issue_column))
        red = PatternFill("solid", fgColor="FCE8E6")
        sheet.conditional_formatting.add("{}2:{}{}".format(letter, letter, sheet.max_row), FormulaRule(formula=['LEN({}2)>0'.format(letter)], fill=red))
        sheet.column_dimensions[letter].width = 48
        for cell in sheet[letter][1:]:
            if cell.value:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                sheet.row_dimensions[cell.row].height = 32
    for column_name in ("播放量", "线索量", "投放消耗"):
        if column_name in [cell.value for cell in sheet[1]] and sheet.max_row > 1:
            letter = get_column_letter(find_column(sheet, column_name))
            sheet.conditional_formatting.add("{}2:{}{}".format(letter, letter, sheet.max_row), CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor="FCE8E6")))
    if sheet.title == "整体数据明细":
        for column_name in NUMBER_FIELDS:
            letter = get_column_letter(find_column(sheet, column_name))
            sheet.conditional_formatting.add("{}2:{}{}".format(letter, letter, sheet.max_row), FormulaRule(formula=["ISBLANK({}2)".format(letter)], fill=PatternFill("solid", fgColor="FCE8E6")))


def find_column(sheet, title):
    for cell in sheet[1]:
        if cell.value == title:
            return cell.column
    raise ValueError("报表字段不存在：{}".format(title))


def summary_value(numerator, denominator):
    if pd.isna(numerator) or pd.isna(denominator) or denominator == 0:
        return float("nan")
    return numerator / denominator


def calculate_overall(frame, labels):
    """整体指标只使用所有必需数值字段齐全的记录，避免未知值充当0。"""
    valid_mask = frame[NUMBER_FIELDS].notna().all(axis=1) & ~frame["数据异常说明"].str.contains("达人名称缺失", regex=False)
    valid = frame.loc[valid_mask]
    total_cost = valid["投放消耗"].sum(min_count=1)
    total_views = valid["播放量"].sum(min_count=1)
    total_leads = valid["线索量"].sum(min_count=1)
    total_sales = valid["成交量"].sum(min_count=1)
    total_gmv = valid["GMV"].sum(min_count=1)
    total_interactions = (valid["点赞数"] + valid["评论数"] + valid["转发数"]).sum(min_count=1)
    groups = [item[0] for item in labels]
    return {
        "总投放消耗": total_cost,
        "总播放量": total_views,
        "总线索量": total_leads,
        "总成交量": total_sales,
        "总GMV": total_gmv,
        "总互动数": total_interactions,
        "整体ROI": summary_value(total_gmv, total_cost),
        "整体CPM": summary_value(total_cost, total_views) * 1000,
        "整体CPL": summary_value(total_cost, total_leads),
        "整体互动率": summary_value(total_interactions, total_views),
        "整体线索转化率": summary_value(total_sales, total_leads),
        "建议复投数量": groups.count("建议复投"),
        "继续观察数量": groups.count("继续观察"),
        "建议淘汰数量": groups.count("建议淘汰"),
        "达人算术平均CPL": valid["CPL"].mean(),
        "达人算术平均互动率": valid["互动率"].mean(),
        "有效汇总记录数": len(valid),
        "待补数据记录数": frame.attrs["quality"]["pending_rows"],
        "异常记录数": (frame["数据异常说明"] != "").sum(),
    }


def find_summary_row(sheet, title):
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == title:
            return row
    raise ValueError("复盘指标不存在：{}".format(title))


def create_quality_report(sheet, frame, overall):
    field_quality = frame.attrs["quality"]["fields"]
    totals = {kind: sum(field_quality[field][kind] for field in FIELDS) for kind in QUALITY_KINDS}
    rows = [
        ("质量指标", "数量"),
        ("总记录数", len(frame)),
        ("有效记录数", overall["有效汇总记录数"]),
        ("待补数据记录数", overall["待补数据记录数"]),
        ("缺失值数量", totals["缺失值"]),
        ("格式异常数量", totals["格式异常"]),
        ("负数异常数量", totals["负数异常"]),
        ("真实0记录数量", frame.attrs["quality"]["zero_rows"]),
        ("异常记录数", overall["异常记录数"]),
    ]
    for row in rows:
        sheet.append(row)
    style_sheet(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = None
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 18
    for row in range(2, 10):
        sheet.cell(row, 2).number_format = COUNT_FORMAT
    header_row = 11
    for column, title in enumerate(("输入字段", "缺失值", "格式异常", "负数异常", "真实0", "需补数据合计"), start=1):
        cell = sheet.cell(header_row, column, title)
        cell.fill = PatternFill("solid", fgColor="233B53")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[header_row].height = 24
    for row_number, field in enumerate(FIELDS, start=header_row + 1):
        counts = field_quality[field]
        values = [field, counts["缺失值"], counts["格式异常"], counts["负数异常"], counts["真实0"], counts["缺失值"] + counts["格式异常"] + counts["负数异常"]]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            cell.font = Font(name="Arial", size=10, color="263238")
            cell.alignment = Alignment(vertical="center", horizontal="right" if column > 1 else "left")
            if column > 1:
                cell.number_format = COUNT_FORMAT
    for letter in "CDEF":
        sheet.column_dimensions[letter].width = 18
    sheet["A23"] = "真实0为有效输入；只有分母为0时对应比率无法计算。"
    sheet["A24"] = "缺失/格式/负数按单元格计数；真实0记录按出现过0的达人去重计数。"
    sheet["A25"] = DISCLAIMER
    for row in (23, 24, 25):
        sheet.cell(row, 1).font = Font(name="Arial", size=10, italic=True, color="607D8B")


def create_config_report(sheet):
    sheet.append(("演示参数", "当前阈值", "分层规则中的作用"))
    settings = [
        ("REINVEST_MIN_ROI", config.REINVEST_MIN_ROI, "ROI达到该值可进入复投判断"),
        ("REINVEST_MAX_CPL", config.REINVEST_MAX_CPL, "CPL不高于该值可进入复投判断，单位：元"),
        ("REINVEST_MIN_LEADS", config.REINVEST_MIN_LEADS, "线索量达到该值可进入复投判断"),
        ("ELIMINATE_BELOW_ROI", config.ELIMINATE_BELOW_ROI, "ROI低于该值触发淘汰判断"),
        ("ELIMINATE_ABOVE_CPL", config.ELIMINATE_ABOVE_CPL, "CPL高于该值触发淘汰判断，单位：元"),
    ]
    for setting in settings:
        sheet.append(setting)
    style_sheet(sheet)
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 58
    for row in (2, 5):
        sheet.cell(row, 2).number_format = RATIO_FORMAT
    for row in (3, 6):
        sheet.cell(row, 2).number_format = MONEY_FORMAT
    sheet["A8"] = "本页阈值仅用于个人项目功能演示，可根据实际业务调整，不代表行业标准。"
    sheet["A9"] = "修改 src/config.py 后重新运行脚本；此页记录生成报告时使用的阈值。"
    sheet["A10"] = DISCLAIMER
    for row in (8, 9, 10):
        sheet.cell(row, 1).font = Font(name="Arial", size=10, italic=True, color="607D8B")


def create_dashboard(sheet, summary, ranking, detail, values, source_count):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A3"
    sheet["A1"] = "达人投放数据看板"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="233B53")
    sheet.row_dimensions[1].height = 30
    sheet["A2"] = "模拟测试数据 · V1.5"
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="607D8B")
    sheet["A3"] = "指标"
    sheet["B3"] = "数值"
    for cell in sheet[3][:2]:
        cell.fill = PatternFill("solid", fgColor="233B53")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[3].height = 24
    kpis = [
        "总投放消耗", "总播放量", "总线索量", "总成交量", "总GMV",
        "整体ROI", "整体CPM", "整体CPL", "整体互动率", "整体线索转化率",
        "建议复投数量", "继续观察数量", "建议淘汰数量",
    ]
    for row, label in enumerate(kpis, start=4):
        sheet.cell(row, 1, label)
        cell = sheet.cell(row, 2, excel_value(values[label]))
        sheet.cell(row, 1).font = Font(name="Arial", size=10, color="263238")
        cell.font = Font(name="Arial", size=10, bold=True, color="233B53")
        cell.alignment = Alignment(horizontal="right", vertical="center")
        sheet.row_dimensions[row].height = 22
        if label in ("总投放消耗", "总GMV", "整体CPM", "整体CPL"):
            cell.number_format = MONEY_FORMAT
        elif label in ("整体互动率", "整体线索转化率"):
            cell.number_format = PERCENT_FORMAT
        elif label == "整体ROI":
            cell.number_format = RATIO_FORMAT
        else:
            cell.number_format = COUNT_FORMAT
        if row % 2 == 0:
            for column in (1, 2):
                sheet.cell(row, column).fill = PatternFill("solid", fgColor="F3F7FA")
    coverage = (("样本总数", source_count), ("有效记录数", values["有效汇总记录数"]), ("待补数据数", values["待补数据记录数"]), ("异常记录数", values["异常记录数"]))
    for row, (label, value) in enumerate(coverage, start=18):
        sheet.cell(row, 1, label)
        cell = sheet.cell(row, 2, value)
        cell.number_format = COUNT_FORMAT
        cell.alignment = Alignment(horizontal="right")
        cell.font = Font(name="Arial", size=10, bold=True, color="233B53")
        sheet.cell(row, 1).font = Font(name="Arial", size=10, color="607D8B")
    sheet["A23"] = "整体指标仅统计数值字段齐全的记录。"
    sheet["A24"] = "缺失和异常不作0；分层规则仅用于演示。"
    sheet["A26"] = "模拟测试数据"
    for address in ("A23", "A24"):
        sheet[address].font = Font(name="Arial", size=10, italic=True, color="607D8B")
    sheet["A26"].font = Font(name="Arial", size=10, bold=True, color="9A6700")
    sheet.column_dimensions["A"].width = 23
    sheet.column_dimensions["B"].width = 23
    sheet.column_dimensions["C"].width = 3
    for column in range(4, 22):
        sheet.column_dimensions[get_column_letter(column)].width = 10

    segment_chart = BarChart()
    segment_chart.type = "col"
    segment_chart.title = "达人分层数量"
    segment_chart.y_axis.title = "达人数量"
    segment_chart.legend = None
    segment_chart.width = 15
    segment_chart.height = 8
    first = find_summary_row(summary, "建议复投数量")
    segment_chart.add_data(Reference(summary, min_col=2, min_row=first, max_row=first + 2))
    segment_chart.set_categories(Reference(summary, min_col=1, min_row=first, max_row=first + 2))
    segment_chart.series[0].graphicalProperties.solidFill = "4B8E86"
    sheet.add_chart(segment_chart, "D3")

    top_chart = BarChart()
    top_chart.type = "bar"
    top_chart.title = "ROI Top10达人"
    top_chart.x_axis.title = "ROI"
    top_chart.legend = None
    top_chart.width = 15
    top_chart.height = 10
    top_chart.x_axis.scaling.orientation = "maxMin"
    top_chart.y_axis.scaling.orientation = "minMax"
    top_chart.y_axis.scaling.min = 0
    top_chart.y_axis.numFmt = RATIO_FORMAT
    roi_column = find_column(ranking, "ROI")
    valid_rank_count = sum(ranking.cell(row, find_column(ranking, "ROI排名")).value is not None for row in range(2, ranking.max_row + 1))
    top_end = min(1 + valid_rank_count, 11)
    if top_end >= 2:
        top_chart.add_data(Reference(ranking, min_col=roi_column, min_row=1, max_row=top_end), titles_from_data=True)
        top_chart.set_categories(Reference(ranking, min_col=1, min_row=2, max_row=top_end))
        top_chart.series[0].graphicalProperties.solidFill = "5B86A6"
    sheet.add_chart(top_chart, "D18")

    scatter = ScatterChart()
    scatter.title = "投放消耗 vs GMV"
    scatter.x_axis.title = "投放消耗（千元）"
    scatter.y_axis.title = "GMV（千元）"
    scatter.x_axis.numFmt = "#,##0,"
    scatter.y_axis.numFmt = "#,##0,"
    scatter.x_axis.scaling.min = 0
    scatter.y_axis.scaling.min = 0
    scatter.legend = None
    scatter.width = 15
    scatter.height = 10
    costs = Reference(detail, min_col=find_column(detail, "投放消耗"), min_row=2, max_row=detail.max_row)
    gmvs = Reference(detail, min_col=find_column(detail, "GMV"), min_row=2, max_row=detail.max_row)
    points = Series(gmvs, costs, title="模拟达人")
    points.marker.symbol = "circle"
    points.marker.size = 4
    points.graphicalProperties.line.noFill = True
    scatter.series.append(points)
    sheet.add_chart(scatter, "M3")


def create_report(frame, output_path):
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "数据看板"
    labels = frame.apply(classify, axis=1)
    overall = calculate_overall(frame, labels)

    detail = workbook.create_sheet("整体数据明细")
    append_frame(detail, frame)
    style_sheet(detail, ["互动率", "线索转化率"], ["投放消耗", "GMV", "CPM", "CPL"], list(COUNT_FIELDS), "ROI", "数据异常说明")
    if "数据声明" in frame.columns:
        declaration_letter = get_column_letter(find_column(detail, "数据声明"))
        detail.column_dimensions[declaration_letter].width = 48
        for cell in detail[declaration_letter][1:]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            detail.row_dimensions[cell.row].height = 32

    ranking = workbook.create_sheet("达人表现排名")
    ranked = frame[["达人名称", "ROI", "CPL", "互动率", "投放消耗", "播放量", "线索量", "成交量", "GMV", "数据异常说明"]].copy()
    complete = frame[NUMBER_FIELDS].notna().all(axis=1) & ~frame["数据异常说明"].str.contains("达人名称缺失", regex=False)
    ranked.insert(1, "ROI排名", frame["ROI"].where(complete).rank(method="min", ascending=False))
    ranked.insert(2, "CPL排名", frame["CPL"].where(complete).rank(method="min", ascending=True))
    ranked.insert(3, "互动率排名", frame["互动率"].where(complete).rank(method="min", ascending=False))
    ranked = ranked.sort_values(["ROI排名", "CPL排名", "互动率排名"], ascending=[True, True, True], na_position="last")
    append_frame(ranking, ranked)
    style_sheet(ranking, ["互动率"], ["投放消耗", "GMV", "CPL"], ["ROI排名", "CPL排名", "互动率排名", "播放量", "线索量", "成交量"], "ROI", "数据异常说明")

    segment = workbook.create_sheet("达人分层建议")
    layered = frame[["达人名称", "ROI", "CPL", "互动率", "线索量", "成交量", "投放消耗", "数据异常说明"]].copy()
    layered.insert(1, "分层结果", [item[0] for item in labels])
    layered.insert(2, "分层原因", [item[1] for item in labels])
    layered.insert(3, "触发规则", [item[2] for item in labels])
    layered["显示顺序"] = layered["分层结果"].map({"建议复投": 0, "继续观察": 1, "建议淘汰": 2})
    layered = layered.sort_values(["显示顺序", "ROI"], ascending=[True, False], na_position="last").drop(columns="显示顺序")
    append_frame(segment, layered)
    style_sheet(segment, ["互动率"], ["CPL", "投放消耗"], ["线索量", "成交量"], "ROI", "数据异常说明")
    segment.column_dimensions[get_column_letter(find_column(segment, "分层原因"))].width = 42
    segment.column_dimensions[get_column_letter(find_column(segment, "触发规则"))].width = 48
    segment["A{}".format(segment.max_row + 2)] = "本页分层仅为演示规则；阈值可在 src/config.py 修改，不代表真实公司或行业标准。"

    quality_report = workbook.create_sheet("数据质量报告")
    create_quality_report(quality_report, frame, overall)

    summary = workbook.create_sheet("整体复盘")
    summary_rows = [("指标", "数值")]
    summary_rows.extend((name, value) for name, value in overall.items())
    summary_rows.extend([
        ("数据声明", DISCLAIMER),
        ("整体口径", "整体指标仅统计达人名称及全部必需数值字段齐全的记录；真实0保留，缺失或异常不作0。"),
        ("平均口径", "达人算术平均CPL和互动率按有效记录中可计算的达人取算术平均。"),
        ("分层规则", "个人项目演示规则；阈值见 src/config.py，不代表真实公司策略或行业标准。"),
    ])
    for row in summary_rows:
        summary.append([excel_value(value) for value in row])
    style_sheet(summary)
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 95
    for row in range(2, summary.max_row + 1):
        label = summary.cell(row, 1).value
        cell = summary.cell(row, 2)
        if label in ("总投放消耗", "总GMV", "整体CPM", "整体CPL", "达人算术平均CPL"):
            cell.number_format = MONEY_FORMAT
        elif label in ("整体互动率", "整体线索转化率", "达人算术平均互动率"):
            cell.number_format = PERCENT_FORMAT
        elif label == "整体ROI":
            cell.number_format = RATIO_FORMAT
        elif isinstance(cell.value, (int, float)):
            cell.number_format = COUNT_FORMAT

    settings = workbook.create_sheet("参数配置说明")
    create_config_report(settings)
    create_dashboard(dashboard, summary, ranking, detail, overall, len(frame))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return overall, labels


def main():
    parser = argparse.ArgumentParser(description=DISCLAIMER)
    parser.add_argument("--generate-sample", action="store_true", help="先生成100条模拟测试数据，再分析该文件")
    parser.add_argument("--input", type=Path, default=ROOT / "data" / "模拟达人投放数据_V1_5.csv", help="输入 CSV 或 XLSX 文件")
    parser.add_argument("--output", type=Path, default=ROOT / "output" / "达人投放复盘报告_V1_5.xlsx", help="输出 XLSX 文件")
    parser.add_argument("--dashboard-image", type=Path, default=ROOT / "output" / "dashboard.png", help="输出16:9看板图片")
    args = parser.parse_args()
    if args.generate_sample:
        if args.input.suffix.lower() != ".csv":
            raise ValueError("--generate-sample 的输入路径必须是 .csv 文件。")
        count = generate_sample(args.input)
        print("已生成 {} 条模拟测试数据：{}".format(count, args.input))
    frame = analyze(load_input(args.input))
    overall, labels = create_report(frame, args.output)
    save_dashboard_image(frame, overall, labels, args.dashboard_image)
    for label, value in (
        ("总记录数", len(frame)),
        ("有效记录", overall["有效汇总记录数"]),
        ("待补数据", overall["待补数据记录数"]),
        ("异常记录", overall["异常记录数"]),
        ("复投数量", overall["建议复投数量"]),
        ("观察数量", overall["继续观察数量"]),
        ("淘汰数量", overall["建议淘汰数量"]),
    ):
        print("{}：{}".format(label, value))
    print("最终报告路径：{}".format(args.output.resolve()))
    print("看板图片路径：{}".format(args.dashboard_image.resolve()))
    print(DISCLAIMER)


if __name__ == "__main__":
    try:
        main()
    except ValueError as error:
        raise SystemExit("错误：{}".format(error))
